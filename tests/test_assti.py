"""Fixtures inteiramente fictícias; não certificam layouts reais da Claro."""
import os
if os.name != 'nt':
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from decimal import Decimal
from unittest.mock import patch
import pymupdf
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from core import read_invoice, export_invoices, InvoiceError, category
from main import MainWindow
from pdf_core import (PDFSource, check_digit, validate_key, extract_text,
                      process_batch, export_results, summarize, sources_from_paths)


def key(number=1):
    body = '35' + '2609' + '11222333000181' + '62' + '001' + f'{number:09}' + '1' + '0' + '0000001'
    assert len(body) == 43
    return body + check_digit(body)


def make_pdf(path, text):
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((30, 40), text, fontsize=9)
        doc.save(path)
    return path


def invoice_xml(value=None):
    k = key()
    return f'''<nfcomProc xmlns="http://www.portalfiscal.inf.br/nfcom">
<NFCom><infNFCom Id="NFCom{k}"><ide><nNF>1</nNF></ide>
<assinante><iCodAssinante>MM001234</iCodAssinante></assinante>
<gFat><CompetFat>202609</CompetFat><dVencFat>2026-10-20</dVencFat><codBarras>{'1'*44}</codBarras></gFat>
<total><vNF>{value or '110.00'}</vNF></total>
<det nItem="1"><prod><xProd>Telefonia</xProd><vProd>80.00</vProd></prod></det>
<det nItem="2"><prod><xProd>Serviço de Gestão</xProd><vProd>20.00</vProd></prod></det>
<det nItem="3"><prod><xProd>Juros</xProd><vProd>2.00</vProd></prod></det>
<det nItem="4"><prod><xProd>Multa</xProd><vProd>3.00</vProd></prod></det>
<det nItem="5"><prod><xProd>Gestão e Autosserviço</xProd><vProd>5.00</vProd></prod></det>
</infNFCom></NFCom><protNFCom><infProt><cStat>100</cStat><chNFCom>{k}</chNFCom></infProt></protNFCom></nfcomProc>'''


class PDFTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_key_structure_and_digit(self):
        self.assertTrue(validate_key(key())[0])
        self.assertEqual(key(), '35260911222333000181620010000000011000000012')  # soma ponderada 350; resto 9; DV 2
        invalid = key()[:-1] + '0'
        self.assertIn('verificador', validate_key(invalid)[1])
        for offset, value in ((0, '00'), (4, '13'), (20, '55'), (6, '11111111111111'), (34, '9')):
            changed = key()[:offset] + value + key()[offset+len(value):]
            self.assertFalse(validate_key(changed)[0])
        self.assertFalse(validate_key(key()[:-1])[0])
        self.assertFalse(validate_key('９'*44)[0])

    def test_spaces_separators_newlines_multiple(self):
        k = key()
        formatted = ' '.join(k[i:i+4] for i in range(0,44,4))
        formatted = formatted.replace(' ', '. ', 2).replace(' ', '\n', 3)
        account, keys = extract_text('Numero da conta: 001234\nNFCom\nChave de acesso: '+formatted+'\nChave de acesso: '+key(2))
        self.assertEqual(account, '001234')
        self.assertEqual(keys, [k, key(2)])
        self.assertEqual(extract_text('NFCom\nChave de acesso: '+k+'\n'+key(2))[1], [k,key(2)])
        self.assertEqual(extract_text('NFCom\nChave de acesso: '+k[:-1])[1], [k[:-1]])

    def test_avoid_payment_phone_and_unrelated_fiscal_key(self):
        text = 'NFCom\nCodigo de barras: '+ '8'*44 + '\nTelefone: 11999998888\nLinha digitavel: '+ '1'*48
        self.assertEqual(extract_text(text)[1], [])
        self.assertEqual(extract_text('Documento sem secao fiscal\n'+key())[1], [])

    def test_batch_duplicate_invalid_no_key_error_and_excel(self):
        paths = [make_pdf(self.root/'one.pdf', 'NFCom\nConta: 001234\nChave de acesso: '+key()),
                 make_pdf(self.root/'two.pdf', 'NFCom\nChave de acesso: '+key()+'\nChave de acesso: '+key(2)),
                 make_pdf(self.root/'invalid.pdf', 'NFCom\nChave de acesso: '+key()[:-1]+'0'),
                 make_pdf(self.root/'empty.pdf', 'Sem informacao fiscal'), self.root/'broken.pdf']
        paths[-1].write_bytes(b'not a PDF')
        results = process_batch(sources_from_paths(paths))
        counts = list(summarize(results).values())
        self.assertEqual(counts, [5,4,3,1,1,1,1])
        self.assertTrue(results[1].duplicate)
        out = self.root/'out.xlsx'
        export_results(results, out)
        with out.open('rb') as file:
            wb = load_workbook(file)
            self.assertEqual(wb.sheetnames, ['Chaves NFCom','Resumo'])
            self.assertEqual(wb.active['C2'].value, key())
            self.assertEqual(wb.active['C2'].data_type, 's')
            self.assertEqual(wb.active['B2'].value, '001234')
            self.assertEqual(wb['Resumo']['B2'].value, 5)
            wb.close()

    def test_198_pdfs_zip_and_folder(self):
        sample = make_pdf(self.root/'sample.pdf','NFCom\nChave de acesso: '+key())
        archive = self.root/'batch.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            for i in range(198):
                z.write(sample, f'invoices/{i}.pdf')
            z.writestr('../ignore.txt', 'irrelevante')
        sources = sources_from_paths([archive])
        self.assertEqual(len(sources),198)
        calls = []
        results = process_batch(sources, lambda n, _: calls.append(n))
        self.assertEqual(len(results),198)
        self.assertEqual(calls[-1],198)
        self.assertEqual(sum(r.duplicate for r in results),197)
        self.assertEqual(sources_from_paths([self.root]),[PDFSource(str(sample))])
        self.assertFalse((self.root/'invoices').exists())

    def test_cancel_preserves_completed(self):
        sample = make_pdf(self.root/'one.pdf','NFCom\nChave de acesso: '+key())
        stop = []
        results = process_batch([PDFSource(str(sample))]*5, lambda *_: stop.append(True), lambda: bool(stop))
        self.assertEqual(len(results),1)

    def test_image_only_and_ocr_unavailable(self):
        sample = make_pdf(self.root/'blank.pdf', '')
        results = process_batch([PDFSource(str(sample))])
        self.assertEqual(results[0].status,'Sem chave')
        self.assertIn('sem texto',results[0].notes)
        with patch('pymupdf.Page.get_textpage_ocr', side_effect=RuntimeError()):
            row = process_batch([PDFSource(str(sample))],ocr=True)[0]
            self.assertIn('OCR indisponível',row.notes)

    def test_ocr_supported_route(self):
        sample = make_pdf(self.root/'blank.pdf', '')
        original_get_text = pymupdf.Page.get_text
        def text_after_ocr(page, *args, **kwargs):
            if kwargs.get('textpage') is not None:
                return 'NFCom\nChave de acesso: '+key()
            return original_get_text(page, *args, **kwargs)
        with patch('pymupdf.Page.get_textpage_ocr',return_value=object()) as ocr, \
             patch('pymupdf.Page.get_text',text_after_ocr):
            row = process_batch([PDFSource(str(sample))],ocr=True)[0]
            self.assertTrue(row.valid)
            self.assertIn('OCR utilizado',row.notes)
            ocr.assert_called_once_with(language='por+eng', dpi=200, full=True)

    @unittest.skipUnless(os.environ.get('TESSDATA_PREFIX'), 'OCR real opcional: configure TESSDATA_PREFIX com por e eng')
    def test_real_ocr_image_pdf(self):
        original = pymupdf.open()
        page = original.new_page()
        page.insert_text((40,60),'NFCom\nConta: 001234\nChave de acesso: '+key(), fontsize=13)
        image = page.get_pixmap(matrix=pymupdf.Matrix(3,3)).tobytes('png')
        scanned = pymupdf.open(); scanned.new_page().insert_image(page.rect,stream=image)
        path = self.root/'scanned.pdf';scanned.save(path)
        scanned.close();original.close()
        rows = process_batch([PDFSource(str(path))],ocr=True)
        self.assertTrue(rows[0].valid, rows[0].notes)
        self.assertEqual(rows[0].key,key())
        self.assertIn('OCR utilizado',rows[0].notes)

    def test_protected_pdf_and_zip_limits(self):
        dest = self.root/'protected.pdf'
        doc = pymupdf.open(); doc.new_page()
        doc.save(dest,encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password');doc.close()
        self.assertEqual(process_batch([PDFSource(str(dest))])[0].status,'Erro de leitura')
        archive = self.root/'bad.zip';archive.write_bytes(b'invalid')
        with self.assertRaises(zipfile.BadZipFile): sources_from_paths([archive])
        good = self.root/'limit.zip'
        with zipfile.ZipFile(good,'w') as z:z.writestr('file.pdf',b'123')
        with patch('pdf_core.MAX_PDF_BYTES',2):
            with self.assertRaises(ValueError):sources_from_paths([good])


class XMLRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root/'invoice.xml'
        self.path.write_text(invoice_xml(),encoding='utf-8')

    def test_financial_classification_and_export_contract(self):
        inv = read_invoice(self.path)
        self.assertEqual((inv.consumption, inv.management, inv.interest, inv.total),
                         tuple(map(Decimal, ('85.00','20.00','5.00','110.00'))))
        self.assertEqual(inv.account,'001234')
        self.assertEqual(category('Gestão e Autosserviço'),'Consumo')
        out = self.root/'xml.xlsx';export_invoices([inv],out)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames,['Resumo das faturas','Itens da NFCOM','Códigos de pagamento'])
        sheet = wb.active
        self.assertEqual([c.value for c in sheet[1]],['Número da conta','Valor total da NFCOM','Vencimento','Valor de consumo','Serviço de gestão','Juros e multa','Número da NFCOM','Competência'])
        self.assertEqual(sheet['D2'].value,85)
        self.assertEqual(sheet['B2'].number_format,'#,##0.00')
        self.assertFalse(sheet.tables)
        self.assertEqual(wb.worksheets[2]['B2'].value,key())
        wb.close()
        with self.assertRaises(InvoiceError): export_invoices([inv,inv],out)

    def test_divergent_totals_protocol_and_unsafe_xml(self):
        for data in (invoice_xml('120.00'),invoice_xml().replace('<cStat>100','<cStat>101'),
                     '<!DOCTYPE x [<!ENTITY a "danger">]><x>&a;</x>'):
            self.path.write_text(data,encoding='utf-8')
            with self.assertRaises(InvoiceError):read_invoice(self.path)


class UITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.window = MainWindow(self.root/'settings');self.window.show()
        self.addCleanup(self.window.close)

    def wait(self, predicate):
        deadline = time.monotonic()+15
        while predicate() and time.monotonic()<deadline:
            self.app.processEvents(); time.sleep(.005)
        self.app.processEvents()
        self.assertFalse(predicate(),'Processamento não terminou')

    def test_tabs_background_pdf_copy_export_and_xml_regression(self):
        w = self.window; tab = w.pdf_tab
        self.assertEqual(w.tabs.count(),2)
        pdf = make_pdf(self.root/'invoice.pdf','NFCom\nConta: 001234\nChave de acesso: '+key()+'\nChave de acesso: '+key(2))
        with patch.object(QFileDialog,'getOpenFileNames',return_value=([str(pdf)],'')):
            tab.choose_pdfs()
        tab.start();self.wait(lambda: tab.worker is not None)
        self.assertEqual(tab.grid.rowCount(),2)
        tab.grid.selectRow(0);tab.copy_selected()
        self.assertEqual(self.app.clipboard().text(),key())
        tab.grid.selectAll();tab.copy_selected()
        self.assertEqual(self.app.clipboard().text(),key()+'\n'+key(2))
        tab.copy_valid();self.assertEqual(self.app.clipboard().text(),key()+'\n'+key(2))
        out=self.root/'ui.xlsx'
        with patch.object(QFileDialog,'getSaveFileName',return_value=(str(out),'')):
            tab.export()
        self.assertTrue(out.is_file())
        w.tabs.setCurrentIndex(1)
        xml = self.root/'invoice.xml';xml.write_text(invoice_xml(),encoding='utf-8')
        w.import_files([str(xml)]);self.wait(lambda:w.worker is not None)
        self.assertEqual(w.grid.rowCount(),1)
        self.assertEqual(w.grid.item(0,2).text(),'R$ 85,00')
        self.assertTrue(w.export_btn.isEnabled())
        xml_out = self.root/'xml-ui.xlsx'
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(xml_out),'')), \
             patch.object(QMessageBox, 'information'):
            w.export();self.wait(lambda:w.worker is not None)
        self.assertTrue(xml_out.is_file())
        self.assertEqual(w.history[0]['count'],1)
        w.tabs.setCurrentIndex(0);self.assertEqual(tab.grid.rowCount(),2)
        w.tabs.setCurrentIndex(1);self.assertEqual(w.grid.rowCount(),1)
        tab.clear();self.assertEqual(tab.grid.rowCount(),0)
        self.assertEqual(w.grid.rowCount(),1)

    def test_folder_zip_import_and_clear(self):
        tab=self.window.pdf_tab
        pdf=make_pdf(self.root/'one.pdf','NFCom\nChave de acesso: '+key())
        with patch.object(QFileDialog,'getExistingDirectory',return_value=str(self.root)):
            tab.choose_folder()
        self.assertEqual(len(tab.sources),1)
        archive=self.root/'one.zip'
        with zipfile.ZipFile(archive,'w') as z:z.write(pdf,'one.pdf')
        with patch.object(QFileDialog,'getOpenFileNames',return_value=([str(archive)],'')):
            tab.choose_zip()
        self.assertEqual(len(tab.sources),2)
        tab.clear();self.assertEqual(len(tab.sources),0)

    def test_close_blocked_until_pdf_cancel_finishes(self):
        tab=self.window.pdf_tab
        pdf=make_pdf(self.root/'one.pdf','NFCom\nChave de acesso: '+key())
        tab.add_paths([pdf]);tab.start()
        with patch.object(self.window,'error') as error:
            self.assertFalse(self.window.close());error.assert_called_once()
        tab.cancel();self.wait(lambda:tab.worker is not None)
        self.assertTrue(self.window.close())


if __name__ == '__main__':
    unittest.main()
