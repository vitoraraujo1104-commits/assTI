"""Leitura local de NFCom e exportação XLSX, sem automação do Excel."""
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import os
import re
import tempfile
import unicodedata
from defusedxml import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

NS = {'n': 'http://www.portalfiscal.inf.br/nfcom'}
ZERO = Decimal('0.00')

class InvoiceError(ValueError):
    pass

@dataclass(frozen=True)
class Item:
    number: str
    description: str
    category: str
    value: Decimal

@dataclass(frozen=True)
class Invoice:
    account: str
    number: str
    period: str
    due: date
    total: Decimal
    consumption: Decimal
    management: Decimal
    interest: Decimal
    key: str
    barcode: str
    items: tuple[Item, ...]
    filename: str


def category(description):
    text = ''.join(c for c in unicodedata.normalize('NFD', description)
                   if unicodedata.category(c) != 'Mn').lower()
    text = re.sub(r'\s+', ' ', text).strip()
    if re.search(r'\bservico de gestao\b', text):
        return 'Serviço de gestão'
    if re.search(r'\b(juros|multa)\b', text):
        return 'Juros e multa'
    return 'Consumo'


def brl(value):
    return 'R$ ' + f'{value:,.2f}'.translate(str.maketrans(',.', '.,'))


def text(node, path, required=True):
    found = node.find(path, NS)
    value = '' if found is None else (found.text or '').strip()
    if required and not value:
        raise InvoiceError(f'Campo obrigatório ausente: {path.replace("n:", "")}')
    return value


def money(node, path):
    raw = text(node, path)
    if not re.fullmatch(r'-?\d+(?:\.\d{1,2})?', raw):
        raise InvoiceError(f'Valor monetário inválido em {path.replace("n:", "")}')
    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise InvoiceError('Valor monetário inválido.') from exc


def read_invoice(path):
    path = Path(path)
    if path.suffix.lower() != '.xml':
        raise InvoiceError('Selecione um arquivo XML.')
    if path.stat().st_size > 25 * 1024 * 1024:
        raise InvoiceError('XML maior que 25 MB. Revise o arquivo.')
    try:
        root = ET.parse(path, forbid_dtd=True, forbid_entities=True, forbid_external=True).getroot()
    except Exception as exc:
        raise InvoiceError('XML inválido ou com estrutura não permitida.') from exc
    infs = root.findall('.//n:infNFCom', NS)
    if len(infs) != 1:
        raise InvoiceError('Esperada uma única NFCOM por arquivo XML.')
    inf = infs[0]
    protocol = root.find('.//n:protNFCom/n:infProt', NS)
    if protocol is None or text(protocol, 'n:cStat') != '100':
        raise InvoiceError('O XML não contém protocolo de autorização com status 100.')
    key = text(protocol, 'n:chNFCom')
    if not re.fullmatch(r'\d{44}', key) or inf.get('Id') != 'NFCom' + key:
        raise InvoiceError('Chave de acesso ausente ou divergente do protocolo.')
    number = text(inf, 'n:ide/n:nNF')
    if not number.isdigit() or int(number) != int(key[25:34]):
        raise InvoiceError('Número da NFCOM diverge da chave de acesso.')
    account = text(inf, 'n:assinante/n:iCodAssinante')
    account = re.sub(r'^MM(?=\d+$)', '', account)
    period_raw = text(inf, 'n:gFat/n:CompetFat')
    due_raw = text(inf, 'n:gFat/n:dVencFat')
    try:
        if not re.fullmatch(r'\d{6}', period_raw) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', due_raw):
            raise ValueError()
        period = datetime.strptime(period_raw, '%Y%m').strftime('%m/%Y')
        due = date.fromisoformat(due_raw)
    except ValueError as exc:
        raise InvoiceError('Competência ou vencimento inválido.') from exc
    total = money(inf, 'n:total/n:vNF')
    barcode = re.sub(r'\s+', '', text(inf, 'n:gFat/n:codBarras'))
    if not barcode.isdigit() or len(barcode) not in (44, 47, 48):
        raise InvoiceError('Código de pagamento com formato inesperado.')
    if len(barcode) == 48 and barcode.startswith('8') and barcode[2] in '68':
        raw = ''.join(barcode[i:i+11] for i in range(0, 48, 12))
        if Decimal(raw[4:15]) / 100 != total:
            raise InvoiceError('Valor do código de pagamento diverge do total da NFCOM.')
    items = []
    totals = {name: ZERO for name in ('Consumo', 'Serviço de gestão', 'Juros e multa')}
    for det in inf.findall('n:det', NS):
        description = text(det, 'n:prod/n:xProd')
        value = money(det, 'n:prod/n:vProd')
        group = category(description)
        totals[group] += value
        items.append(Item(det.get('nItem', ''), description, group, value))
    if not items:
        raise InvoiceError('NFCOM sem itens de cobrança.')
    if sum(totals.values(), ZERO) != total:
        raise InvoiceError(f'Soma dos itens ({brl(sum(totals.values(), ZERO))}) difere do total '
                           f'({brl(total)}). Revise descontos e ajustes no XML.')
    return Invoice(account, number, period, due, total, totals['Consumo'],
                   totals['Serviço de gestão'], totals['Juros e multa'], key,
                   barcode, tuple(items), path.name)


def export_invoices(invoices, destination):
    invoices = list(invoices)
    if not invoices:
        raise InvoiceError('Não há faturas para exportar.')
    if len({i.key for i in invoices}) != len(invoices):
        raise InvoiceError('Existem faturas duplicadas no lote.')
    destination = Path(destination)
    wb = Workbook()
    summary = wb.active
    summary.title = 'Resumo das faturas'
    detail = wb.create_sheet('Itens da NFCOM')
    codes = wb.create_sheet('Códigos de pagamento')
    summary.append(['Número da conta', 'Valor total da NFCOM', 'Vencimento', 'Valor de consumo',
                    'Serviço de gestão', 'Juros e multa', 'Número da NFCOM', 'Competência'])
    detail.append(['Número da conta', 'Número da NFCOM', 'Competência', 'Item', 'Descrição',
                   'Classificação', 'Valor do item', 'Arquivo XML'])
    codes.append(['Número da conta', 'Chave de acesso', 'Código de barras / linha digitável'])
    for i in sorted(invoices, key=lambda i: (i.period[3:] + i.period[:2], i.account, i.number)):
        summary.append([i.account, i.total, i.due, i.consumption, i.management, i.interest, i.number, i.period])
        codes.append([i.account, i.key, i.barcode])
        for p in i.items:
            detail.append([i.account, i.number, i.period, p.number, p.description, p.category, p.value, i.filename])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                # XML text is always literal, including strings starting with '='.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Calibri', size=11, bold=cell.row == 1)
        for column in sheet.columns:
            length = max(len(str(c.value or '')) for c in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(length + 3, 18), 65)
    for row in summary.iter_rows(min_row=2):
        for index in (1, 3, 4, 5):
            row[index].number_format = '#,##0.00'
        row[2].number_format = 'dd/mm/yyyy'
    for row in detail.iter_rows(min_row=2):
        row[6].number_format = '#,##0.00'
    descriptor, temp = tempfile.mkstemp(prefix='.assti-', suffix='.xlsx', dir=destination.parent)
    os.close(descriptor)
    try:
        wb.save(temp)
        os.replace(temp, destination)
    finally:
        wb.close()
        if os.path.exists(temp):
            os.unlink(temp)
